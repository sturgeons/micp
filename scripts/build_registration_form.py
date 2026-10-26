from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "软著信息登记表_已填写_保留表格.docx"


def set_font(run, name: str = "宋体", size: float = 9.0, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def set_cell_text(cell, text: str, size: float = 9.0, bold: bool = False, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def set_cell_margins(table, top=25, start=35, bottom=25, end=35) -> None:
    tbl_pr = table._tbl.tblPr
    mar = tbl_pr.first_child_found_in("w:tblCellMar")
    if mar is None:
        mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(mar)
    for side, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_fixed_layout(table, total_width_cm: float) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(total_width_cm * 567)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)

    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table)
    set_cell_margins(table)
    set_fixed_layout(table, 17.7)

    left_w = 6.3
    right_w = 11.4

    def add_row(left: str, right: str = "", merge: bool = False, size: float = 9.0, bold: bool = False, align=None):
        row = table.add_row()
        set_cell_width(row.cells[0], left_w)
        set_cell_width(row.cells[1], right_w)
        if merge:
            cell = row.cells[0].merge(row.cells[1])
            set_cell_text(cell, left, size=size, bold=bold, align=align)
        else:
            set_cell_text(row.cells[0], left, size=size, bold=bold)
            set_cell_text(row.cells[1], right, size=size)

    add_row("计算机著作权项目确认单", merge=True, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_row("（本确认单作为合同的一部分；如软件信息内容与合同有差异，以此确认单为准）", merge=True, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_row("注意事项：此表一经提交无法修改，请谨慎填写！", merge=True, size=9, bold=True)
    add_row("合同编号：", "待填写")
    add_row("申请软件著作权是否用于双软认证：", "□是 ☑否")
    add_row("软件全称：", "MICP-椰壳纤维协同固化砂土性能分析与预测系统")
    add_row("软件简称：（可无）", "MICP 固化砂土分析预测系统")
    add_row("版本号：", "V1.0")
    add_row("软件分类：", "☑应用软件  □嵌入式软件  □中间件  □操作系统")
    add_row("开发方式：", "☑独立开发  □合作开发  □委托开发  □下达任务开发")
    add_row("开发目的：（20字以内）", "MICP固化砂土性能分析预测")
    add_row("面向领域/行业：", "岩土工程、土体固化、MICP 试验研究")
    add_row("软件开发完成日期：", "2026 年 6 月 23 日")
    add_row("软件是否发表", "□是  发表的日期：/ 年 / 月 / 日\n发表的地点：/\n☑否")
    add_row("著作权人信息（个人/公司名称）", "待填写")
    add_row("证件号码（身份证/公司营业执照注册号）", "待填写")
    add_row("软件开发的硬件环境：", "通用 PC 或服务器，CPU 2 核及以上，内存 8GB 及以上。")
    add_row("软件运行的硬件环境", "通用 PC、笔记本电脑或云服务器，支持浏览器访问。")
    add_row("开发该软件的操作系统：", "macOS、Windows 或 Linux。")
    add_row("软件开发工具/环境：", "Python、Flask、SQLite、scikit-learn、HTML5、CSS3、JavaScript。")
    add_row("软件的运行平台/操作系统：", "Windows 10 及以上、macOS、Linux。")
    add_row("软件运行支撑环境/支持软件：", "Python 3.x、Flask、SQLite、scikit-learn、Chrome/Edge/Safari。")
    add_row("编程语言及版本号：", "Python 3.x、JavaScript ES6、HTML5、CSS3。")
    add_row("软件代码行数：", "5748 行")
    add_row(
        "软件的主要功能（100字以上200字以下）",
        "本软件面向椰壳纤维协同 MICP 固化砂土试验场景，提供用户与项目管理、试验参数录入、参数模板复用、历史样本管理、模型训练、固化性能预测、敏感性分析、统计分析、图表可视化和报告生成导出等功能，辅助科研与工程人员完成试验数据分析、方案优选和成果归档。",
    )
    add_row(
        "软件的技术特点：（50字以上100字以下）",
        "软件采用 Flask 前后端一体化架构，结合 SQLite 数据存储和 scikit-learn 模型训练能力，支持离线工作台与完整后端服务两种运行方式，可输出多指标预测、统计分析和可视化报告。",
    )
    add_row(
        "软件的类型",
        "□APP  □游戏软件  □教育软件  □金融软件  □医疗软件\n"
        "□地理信息软件  □云计算软件  □信息安全软件  □大数据软件\n"
        "☑人工智能软件  □VR软件  □5G软件  □小程序  □物联网软件  □智能城市软件",
    )
    add_row("申请人：姓名/公司名称", "待填写")
    add_row("详细地址：", "待填写")
    add_row("E-mail：", "待填写")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
