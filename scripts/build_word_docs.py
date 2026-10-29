from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "word"
SCREENSHOT_DIR = DOCS / "截图资产"

MATERIAL_DOC = OUT / "MICP软著材料汇编.docx"
SPEC_DOC = OUT / "MICP设计说明书.docx"
CODE_FULL_DOC = OUT / "MICP源程序全量清单.docx"
CODE_EXCERPT_DOC = OUT / "MICP源程序鉴别材料（前后各30页）.docx"
HEADER_TITLE = "MICP-椰壳纤维协同固化砂土性能分析与预测系统 V1.0"
LINES_PER_PAGE = 50
# 8 磅等宽字体、左右边距 1.4 cm 时，一行大约容纳的 ASCII 列宽；中文按 2 列计
CODE_COLUMNS = 92

MATERIAL_FILES = [
    DOCS / "软著材料清单.md",
    DOCS / "软件基础信息与功能说明.md",
    DOCS / "设计说明书.md",
    DOCS / "流程图.md",
    DOCS / "用户操作手册.md",
    DOCS / "测试报告.md",
    DOCS / "部署说明.md",
    DOCS / "截图说明.md",
]

# 仅收录原创业务代码：算法包、接口层、数据层与前端逻辑；不含 README / 依赖清单 / 样式表
CODE_FILES = [
    ROOT / "backend" / "micp" / "__init__.py",
    ROOT / "backend" / "micp" / "schemas.py",
    ROOT / "backend" / "micp" / "mechanism.py",
    ROOT / "backend" / "micp" / "models.py",
    ROOT / "backend" / "micp" / "evaluation.py",
    ROOT / "backend" / "micp" / "sensitivity.py",
    ROOT / "backend" / "micp" / "optimizer.py",
    ROOT / "backend" / "micp" / "statistics.py",
    ROOT / "backend" / "micp" / "validators.py",
    ROOT / "backend" / "micp" / "synthetic.py",
    ROOT / "backend" / "micp" / "report.py",
    ROOT / "backend" / "db.py",
    ROOT / "backend" / "server.py",
    ROOT / "backend" / "ml.py",
    ROOT / "backend" / "wsgi.py",
    ROOT / "backend" / "bootstrap_admin.py",
    ROOT / "backend" / "static" / "app.js",
    ROOT / "backend" / "templates" / "index.html",
]


def set_east_asia_font(run, font_name: str) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.first_child_found_in("w:tblCellMar")
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for margin, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tbl_cell_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa: int = 9360) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = OxmlElement("w:tblLayout")
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_pr.append(tbl_layout)


def style_document(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    if "Code" not in styles:
        code_style = styles.add_style("Code", 1)
    else:
        code_style = styles["Code"]
    code_style.font.name = "Courier New"
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code_style.font.size = Pt(8)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1.0

    set_header(section, title)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run2 = paragraph.add_run()
    run2._r.append(instr)
    run3 = paragraph.add_run()
    run3._r.append(separate)
    run3._r.append(end)
    for item in (run, run2, run3):
        item.font.size = Pt(9)
        item.font.color.rgb = RGBColor(80, 80, 80)
        set_east_asia_font(item, "宋体")


def set_header(section, title: str) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    paragraph = section.header.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(16.4), WD_TAB_ALIGNMENT.RIGHT)
    run = paragraph.add_run(title)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(80, 80, 80)
    set_east_asia_font(run, "宋体")
    paragraph.add_run("\t")
    add_page_field(paragraph)
    footer = section.footer.paragraphs[0]
    footer.clear()


def set_page_start(section, start: int) -> None:
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def add_cover(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(96)
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string("0B2545")
    set_east_asia_font(run, "黑体")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(36)
    run = p.add_run(subtitle)
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string("555555")
    set_east_asia_font(run, "宋体")

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_width(table, 7600)
    set_cell_margins(table)
    rows = [
        ("软件全称", "MICP-椰壳纤维协同固化砂土性能分析与预测系统"),
        ("软件简称", "MICP 固化砂土分析预测系统"),
        ("版本号", "V1.0"),
        ("材料类型", subtitle),
        ("整理日期", "2026-10-29"),
    ]
    for row, (label, value) in zip(table.rows, rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], "F2F4F7")
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    set_east_asia_font(run, "宋体")
                    run.font.size = Pt(10.5)

    doc.add_page_break()


def add_markdown_table(doc: Document, lines: list[str], start: int) -> int:
    header = [c.strip() for c in lines[start].strip().strip("|").split("|")]
    rows = []
    i = start + 2
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append([c.strip().replace("`", "") for c in lines[i].strip().strip("|").split("|")])
        i += 1

    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    set_cell_margins(table)
    for idx, text in enumerate(header):
        table.rows[0].cells[idx].text = text
        set_cell_shading(table.rows[0].cells[idx], "F2F4F7")
    for row in rows:
        cells = table.add_row().cells
        for idx in range(len(header)):
            cells[idx].text = row[idx] if idx < len(row) else ""
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs:
                    set_east_asia_font(run, "宋体")
                    run.font.size = Pt(9)
    return i


def add_markdown_content(doc: Document, md_path: Path, images: list[Path] | None = None) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    in_code = False
    image_index = 0
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if line.strip().startswith("```mermaid"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                i += 1
            i += 1
            if images and image_index < len(images) and images[image_index].exists():
                picture = doc.add_paragraph()
                picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
                picture.add_run().add_picture(str(images[image_index]), width=Cm(15.8))
                image_index += 1
            continue
        if line.strip().startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph(style="Code")
            p.add_run(raw)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|\s*-+", lines[i + 1].strip()):
            i = add_markdown_table(doc, lines, i)
            continue
        image_line = re.match(r"^!\[(.*?)\]\((.+?)\)$", line.strip())
        if image_line:
            caption, src = image_line.group(1), image_line.group(2)
            image_path = Path(src)
            if not image_path.is_absolute():
                image_path = (md_path.parent / src).resolve()
            if image_path.exists():
                picture = doc.add_paragraph()
                picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
                picture.paragraph_format.space_before = Pt(6)
                picture.paragraph_format.space_after = Pt(2)
                picture.add_run().add_picture(str(image_path), width=Cm(15.5))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.space_before = Pt(0)
                cap.paragraph_format.space_after = Pt(8)
                run = cap.add_run(caption)
                run.bold = True
                run.font.size = Pt(10.5)
                set_east_asia_font(run, "黑体")
            i += 1
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            level = min(len(heading.group(1)), 3)
            doc.add_heading(heading.group(2).replace("`", ""), level=level)
            i += 1
            continue
        bullet = re.match(r"^\s*-\s+(.*)$", line)
        if bullet:
            text = bullet.group(1).replace("`", "")
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(text)
            i += 1
            continue
        numbered = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if numbered:
            text = numbered.group(1).replace("`", "")
            p = doc.add_paragraph(style="List Number")
            p.add_run(text)
            i += 1
            continue
        p = doc.add_paragraph()
        p.add_run(line.replace("`", ""))
        i += 1


def add_screenshots(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("运行截图附件", level=1)
    for image in sorted(SCREENSHOT_DIR.glob("*.png")):
        caption = image.stem
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"图 {caption}")
        r.bold = True
        set_east_asia_font(r, "黑体")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image), width=Cm(15.8))


def build_material_doc() -> None:
    doc = Document()
    style_document(doc, "MICP 软著材料汇编")
    add_cover(doc, "MICP 软著材料汇编", "说明文档与运行截图")
    for index, md in enumerate(MATERIAL_FILES):
        if index:
            doc.add_page_break()
        add_markdown_content(doc, md)
    add_screenshots(doc)
    doc.save(MATERIAL_DOC)


def display_columns(text: str) -> int:
    width = 0
    for char in text:
        width += 2 if ord(char) > 127 else 1
    return width


def wrap_source_line(text: str) -> list[str]:
    """按显示列宽拆行，保证写进 Word 后一行不再自动换行。"""
    if display_columns(text) <= CODE_COLUMNS:
        return [text]
    parts: list[str] = []
    buf = ""
    width = 0
    for char in text:
        char_width = 2 if ord(char) > 127 else 1
        if buf and width + char_width > CODE_COLUMNS:
            parts.append(buf)
            buf = char
            width = char_width
        else:
            buf += char
            width += char_width
    if buf:
        parts.append(buf)
    return parts


def collect_source_lines() -> list[str]:
    lines: list[str] = []
    for path in CODE_FILES:
        rel = path.relative_to(ROOT).as_posix()
        lines.append(f"# ===== {rel} =====")
        text = path.read_text(encoding="utf-8", errors="replace")
        if text.endswith("\n"):
            text = text[:-1]
        for source_line in text.split("\n"):
            lines.extend(wrap_source_line(source_line))
    return lines


def paginate(lines: list[str]) -> list[list[str]]:
    return [lines[i:i + LINES_PER_PAGE] for i in range(0, len(lines), LINES_PER_PAGE)]


def apply_code_page(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.4)
    set_header(section, HEADER_TITLE)


def write_code_line(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(13.2)
    p_pr = paragraph._p.get_or_add_pPr()
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    p_pr.append(snap)
    run = paragraph.add_run(text if text else " ")
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(8)
    return paragraph


def write_code_pages(doc: Document, pages: list[list[str]]) -> None:
    for index, page in enumerate(pages):
        for line_index, line in enumerate(page):
            paragraph = write_code_line(doc, line)
            last_of_page = line_index == len(page) - 1 and index < len(pages) - 1
            if last_of_page:
                break_run = paragraph.add_run()
                page_break = OxmlElement("w:br")
                page_break.set(qn("w:type"), "page")
                break_run._r.append(page_break)


def spec_images() -> list[Path]:
    folder = DOCS / "流程图资产"
    return [folder / f"设计说明书-{i:02d}.png" for i in range(1, 6)]


def build_spec_doc() -> None:
    doc = Document()
    style_document(doc, HEADER_TITLE)
    add_cover(doc, "设计说明书", "软件著作权文档鉴别材料")
    add_markdown_content(doc, DOCS / "设计说明书.md", spec_images())
    doc.save(SPEC_DOC)


def build_code_docs() -> tuple[int, int]:
    pages = paginate(collect_source_lines())
    full = Document()
    apply_code_page(full.sections[0])
    write_code_pages(full, pages)
    full.save(CODE_FULL_DOC)

    total = len(pages)
    if total <= 60:
        excerpt_pages = pages
        ranges = [(0, total, 1)]
    else:
        ranges = [(0, 30, 1), (total - 30, total, total - 29)]
    excerpt = Document()
    first = True
    for start, end, page_no in ranges:
        if first:
            section = excerpt.sections[0]
            first = False
        else:
            section = excerpt.add_section(WD_SECTION.NEW_PAGE)
        apply_code_page(section)
        set_page_start(section, page_no)
        write_code_pages(excerpt, pages[start:end])
    excerpt.save(CODE_EXCERPT_DOC)
    return total, sum(len(page) for page in pages)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_spec_doc()
    page_count, line_count = build_code_docs()
    print(SPEC_DOC)
    print(f"{CODE_FULL_DOC}  pages={page_count} lines={line_count}")
    print(CODE_EXCERPT_DOC)


if __name__ == "__main__":
    main()
